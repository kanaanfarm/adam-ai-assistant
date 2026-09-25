"""Pure document/attachment processing services for Adam Acquisition v2.1."""
from __future__ import annotations
import io, mimetypes, re
from pathlib import Path

MAX_ATTACHMENT_BYTES = 20 * 1024 * 1024
MAX_EXTRACTED_CHARS = 120_000
TEXT_EXTS={'.txt','.md','.csv','.log'}; DOC_EXTS={'.pdf','.docx'}; SHEET_EXTS={'.xlsx','.xlsm'}; IMAGE_EXTS={'.png','.jpg','.jpeg','.webp','.gif','.bmp'}

class DocumentProcessingError(RuntimeError): pass

def safe_filename(name, default='Adam_Document'):
    v=re.sub(r'[^A-Za-z0-9._-]+','_',str(name or '').strip()).strip('._')
    return v or default

def classify_attachment(filename,mimetype=''):
    name=str(filename or 'attachment').strip(); ext=Path(name).suffix.lower(); mime=str(mimetype or mimetypes.guess_type(name)[0] or '').lower()
    if ext in DOC_EXTS: return 'document'
    if ext in SHEET_EXTS: return 'spreadsheet'
    if ext in TEXT_EXTS: return 'text'
    if mime.startswith('image/') or ext in IMAGE_EXTS: return 'image'
    raise DocumentProcessingError(f'{name}: unsupported file type.')

def _bounded(text,max_chars=MAX_EXTRACTED_CHARS): return str(text or '')[:max(1,int(max_chars or MAX_EXTRACTED_CHARS))]

def extract_xlsx_text(data,max_chars=MAX_EXTRACTED_CHARS):
    try: from openpyxl import load_workbook
    except Exception as exc: raise DocumentProcessingError('Excel reading requires openpyxl. Run INSTALL_REQUIREMENTS.bat once.') from exc
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True); chunks=[]
    for ws in wb.worksheets[:12]:
        chunks.append(f'\n--- SHEET: {ws.title} ---'); count=0
        for row in ws.iter_rows(values_only=True):
            vals=['' if v is None else str(v) for v in row]
            if any(v.strip() for v in vals): chunks.append(' | '.join(vals)); count+=1
            if count>=500: chunks.append('[Sheet truncated after 500 non-empty rows]'); break
    return _bounded('\n'.join(chunks),max_chars)

def extract_text_bytes(filename,data,max_chars=MAX_EXTRACTED_CHARS):
    name=str(filename or 'attachment').strip(); ext=Path(name).suffix.lower()
    if len(data)>MAX_ATTACHMENT_BYTES: raise DocumentProcessingError(f'{name}: file is larger than 20 MB.')
    if ext=='.pdf':
        try: from pypdf import PdfReader; return _bounded('\n'.join((p.extract_text() or '') for p in PdfReader(io.BytesIO(data)).pages),max_chars)
        except Exception as exc: raise DocumentProcessingError(f'{name}: PDF could not be read.') from exc
    if ext=='.docx':
        try:
            from docx import Document
            return _bounded('\n'.join(p.text for p in Document(io.BytesIO(data)).paragraphs),max_chars)
        except Exception as exc: raise DocumentProcessingError(f'{name}: DOCX could not be read.') from exc
    if ext in SHEET_EXTS: return extract_xlsx_text(data,max_chars)
    if ext in TEXT_EXTS: return _bounded(data.decode('utf-8',errors='replace'),max_chars)
    raise DocumentProcessingError('Supported document types: PDF, DOCX, XLSX/XLSM, TXT, MD, CSV and LOG.')

def process_attachment_bytes(filename,data,mimetype=''):
    name=str(filename or 'attachment').strip()
    if len(data)>MAX_ATTACHMENT_BYTES: raise DocumentProcessingError(f'{name}: file is larger than 20 MB.')
    kind=classify_attachment(name,mimetype)
    if kind=='image': return {'name':name,'kind':'image','mime':str(mimetype or mimetypes.guess_type(name)[0] or 'image/jpeg').lower(),'bytes':data}
    return {'name':name,'kind':kind,'content':extract_text_bytes(name,data)}

def build_analysis_prompt(instruction,extracted_text):
    return 'Analyze only the supplied document. Do not invent missing clauses. Separate document facts from recommendations.\nOWNER INSTRUCTION:\n'+str(instruction or 'Analyze this document.').strip()+'\nDOCUMENT:\n'+_bounded(extracted_text)

def build_docx_bytes(title,body):
    try: from docx import Document
    except Exception as exc: raise DocumentProcessingError('DOCX creation requires python-docx. Run INSTALL_REQUIREMENTS.bat once.') from exc
    doc=Document(); title=str(title or '').strip()
    if title: doc.add_heading(title,level=1)
    for block in str(body or '').split('\n\n'):
        t=block.strip()
        if t: doc.add_paragraph(t)
    stream=io.BytesIO(); doc.save(stream); return stream.getvalue()

def public_processing_capabilities():
    return {'max_attachment_mb':20,'max_extracted_chars':MAX_EXTRACTED_CHARS,'supported_document_types':['pdf','docx'],'supported_spreadsheet_types':['xlsx','xlsm'],'supported_text_types':['txt','md','csv','log'],'supported_image_types':['png','jpg','jpeg','webp','gif','bmp'],'ai_provider_required_for_extraction':False,'network_required_for_extraction':False}
